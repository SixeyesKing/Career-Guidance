import json
import re
from datetime import date, datetime
from pathlib import Path

import yaml
from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
SOURCE_DIR = ROOT / "知识库"
WORKBOOK = SOURCE_DIR / "岗位招聘汇总_300版.xlsx"
OUTPUT_DIR = ROOT / "data" / "knowledge" / "generated"
MANIFEST_PATH = OUTPUT_DIR / ".generated_files.json"
ROLE_MAP = {"后端开发": "backend", "前端开发": "frontend", "测试开发": "testing"}
CATEGORY_DIRS = {
    "01_岗位能力": "job_skill",
    "02_招聘流程": "recruitment_process",
    "03_时间安排": "recruitment_timeline",
    "04_学习路径": "learning_path",
    "05_面试准备": "interview_preparation",
}
GENERATED_PATHS = set()


def scalar_text(value):
    if value is None:
        return ""
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    return str(value).strip()


def slug(value):
    cleaned = re.sub(r"[^\w.-]+", "-", value.strip(), flags=re.UNICODE).strip("-.")
    return cleaned[:70] or "document"


def frontmatter(metadata, content):
    header = yaml.safe_dump(metadata, allow_unicode=True, sort_keys=False).strip()
    return "---\n{0}\n---\n\n{1}\n".format(header, content.strip())


def row_dict(headers, row):
    return {header: value for header, value in zip(headers, row)}


def read_workbook():
    workbook = load_workbook(WORKBOOK, read_only=True, data_only=True)
    ledger_sheet = workbook["岗位台账"]
    raw_sheet = workbook["原始资料"]
    ledger_rows = ledger_sheet.iter_rows(values_only=True)
    raw_rows = raw_sheet.iter_rows(values_only=True)
    ledger_headers = next(ledger_rows)
    raw_headers = next(raw_rows)
    raw_by_sequence = {}
    for row in raw_rows:
        raw = row_dict(raw_headers, row)
        sequence = raw.get("原表序号")
        if sequence is not None:
            raw_by_sequence[str(sequence).strip()] = raw
    records = []
    for row in ledger_rows:
        item = row_dict(ledger_headers, row)
        if not any(value is not None and str(value).strip() for value in item.values()):
            continue
        role = ROLE_MAP.get(scalar_text(item.get("岗位方向")))
        sequence = scalar_text(item.get("序号"))
        if role and sequence:
            records.append((item, raw_by_sequence.get(sequence, {})))
    workbook.close()
    return records


def review_value(value):
    return {
        "已核实": "verified",
        "待复核": "pending_review",
    }.get(scalar_text(value), "not_reviewed")


def posting_content(ledger, raw):
    fields = [
        ("公司", ledger.get("公司")),
        ("岗位方向", ledger.get("岗位方向")),
        ("岗位名称", ledger.get("岗位名称")),
        ("招聘类型", ledger.get("招聘类型")),
        ("面向人群/经验", ledger.get("面向人群/经验")),
        ("工作地点", ledger.get("工作地点")),
        ("职责速览", ledger.get("职责速览")),
        ("要求速览", ledger.get("要求速览")),
        ("与项目的关系", ledger.get("与项目的关系")),
    ]
    lines = ["# 招聘岗位记录"]
    lines.extend("\n**{0}：** {1}".format(label, scalar_text(value)) for label, value in fields if scalar_text(value))
    raw_text = scalar_text(raw.get("原始岗位内容"))
    if raw_text:
        lines.extend(["", "## 原始岗位描述", "", raw_text])
    lines.extend([
        "",
        "## 资料说明",
        "",
        "复核状态沿用原始岗位台账标记。‘已核实’表示台账中的整理状态，不代表本程序再次访问招聘网站独立核验。请以来源页面当前内容为准。",
    ])
    return "\n".join(lines)


def write_postings():
    records = read_workbook()
    for ledger, raw in records:
        sequence = int(ledger["序号"])
        role = ROLE_MAP[scalar_text(ledger.get("岗位方向"))]
        company = scalar_text(ledger.get("公司")) or "公司未提供"
        title = scalar_text(ledger.get("岗位名称")) or "岗位名称未提供"
        url = scalar_text(ledger.get("来源链接")) or scalar_text(raw.get("来源链接"))
        metadata = {
            "id": "job-posting-{0:03d}".format(sequence),
            "title": "{0} - {1}".format(company, title),
            "job_category": role,
            "document_type": "job_posting",
            "source_name": company,
            "source_url": url,
            "updated_at": scalar_text(ledger.get("页面更新时间")) or "未提供",
            "review_status": review_value(ledger.get("复核状态")),
            "source_granularity": "homepage" if url and re.fullmatch(r"https?://[^/]+/?", url) else "posting_detail",
            "is_placeholder": False,
        }
        output = OUTPUT_DIR / "postings" / role / "{0:03d}-{1}.md".format(sequence, slug(company))
        output.parent.mkdir(parents=True, exist_ok=True)
        output.write_text(frontmatter(metadata, posting_content(ledger, raw)), encoding="utf-8")
        GENERATED_PATHS.add(output.relative_to(OUTPUT_DIR).as_posix())
    return len(records)


def markdown_heading(text, label):
    match = re.search(r"^##\s*" + re.escape(label) + r"\s*\n([^\n]+)", text, re.MULTILINE)
    return match.group(1).strip() if match else ""


def applicable_roles(text):
    listed_roles = [name for name in ROLE_MAP if name in markdown_heading(text, "岗位类别")]
    if listed_roles:
        return [(name, ROLE_MAP[name]) for name in listed_roles]
    if "所有岗位" in text or "通用" in markdown_heading(text, "岗位类别"):
        return list(ROLE_MAP.items())
    return []


def document_review(text):
    if "整理时间" in text or "来源信息" in text:
        return "not_reviewed"
    return "not_reviewed"


def write_general_documents():
    written = 0
    for folder, document_type in CATEGORY_DIRS.items():
        source_folder = SOURCE_DIR / folder
        if not source_folder.exists():
            continue
        for source_path in sorted(source_folder.glob("*.md")):
            if source_path.name.lower() == "readme.md":
                continue
            text = source_path.read_text(encoding="utf-8-sig").strip()
            if not text:
                continue
            roles = applicable_roles(text)
            if not roles:
                continue
            source_match = re.search(r"\*\*来源名称\*\*\s*[：:]\s*([^\n]+)", text)
            source_name = source_match.group(1).strip() if source_match else "团队整理（来源待核验）"
            links = re.findall(r"https?://[^\s)]+", text)
            source_url = links[0].rstrip(".,;，。；") if links else ""
            is_time_sensitive = folder == "03_时间安排"
            is_template = bool(re.search(r"XX|\[待补充\]|占位|模板一|模板二|模板三", text))
            updated_match = re.search(r"(?:整理时间|更新时间|发布日期)\s*[：:]\s*(\d{4}(?:年\d{1,2}月\d{1,2}日|[-/]\d{1,2}(?:[-/]\d{1,2})?)?)", text)
            updated_at = updated_match.group(1) if updated_match else "未提供"
            for role_name, role_id in roles:
                source_id = "guide-{0}-{1}".format(slug(source_path.stem), role_id)
                metadata = {
                    "id": source_id,
                    "title": source_path.stem,
                    "job_category": role_id,
                    "document_type": document_type,
                    "source_name": source_name,
                    "source_url": source_url,
                    "updated_at": updated_at,
                    "review_status": document_review(text),
                    "freshness": "time_sensitive" if is_time_sensitive else "general",
                    "content_kind": "template" if is_template else "guide",
                    "is_placeholder": False,
                }
                body = "适用岗位：{0}\n\n{1}".format(role_name, text)
                output = OUTPUT_DIR / "guides" / role_id / "{0}.md".format(source_id)
                output.parent.mkdir(parents=True, exist_ok=True)
                output.write_text(frontmatter(metadata, body), encoding="utf-8")
                GENERATED_PATHS.add(output.relative_to(OUTPUT_DIR).as_posix())
                written += 1
    return written


def main():
    if not WORKBOOK.exists():
        raise FileNotFoundError("岗位工作簿不存在：{0}".format(WORKBOOK))
    previous_paths = set()
    if MANIFEST_PATH.exists():
        try:
            previous_paths = set(json.loads(MANIFEST_PATH.read_text(encoding="utf-8")))
        except (ValueError, TypeError):
            previous_paths = set()
    posting_count = write_postings()
    guide_count = write_general_documents()
    for relative_path in previous_paths - GENERATED_PATHS:
        target = (OUTPUT_DIR / relative_path).resolve()
        try:
            target.relative_to(OUTPUT_DIR.resolve())
        except ValueError:
            continue
        if target.is_file() and target.suffix.lower() == ".md":
            target.unlink()
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    MANIFEST_PATH.write_text(
        json.dumps(sorted(GENERATED_PATHS), ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )
    print("生成招聘岗位资料：{0}".format(posting_count))
    print("生成通用知识资料：{0}".format(guide_count))
    print("输出目录：{0}".format(OUTPUT_DIR))


if __name__ == "__main__":
    main()
